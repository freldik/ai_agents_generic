"""Aktualizuje sloupce Investovano a Hodnota v sporeni.xlsx podle exportů z brokerů.

Exporty se hledají ve stromu:

    <kořen>/<Vlastnik>/<Broker>/<MENA>_....xlsx

např. `D:\\Work\\Finance\\Test\\Marek\\XTB\\USD_51959804_2006-01-01_2026-08-13.xlsx`.
V každé složce brokera smí být pro jednu měnu **právě jeden** soubor — konkrétní název
se proto nikde nezadává. Měna se pozná z předpony názvu (`USD_`, `EUR_`, ...).

Na listu akcií je jeden řádek = jedna pozice (Vlastnik, Broker, Mena). Export
`<Vlastnik>/<Broker>/USD_*.xlsx` plní jen řádky s týmž vlastníkem, brokerem a
`Mena=USD` — sešit čísla přepočítává kurzem podle sloupce Mena, takže USD částka
v EUR řádku by dala špatné CZK.

Páruje se podle tickeru (první část sloupce `Nazev akcie` před čárkou), u řádků bez
tickeru podle názvu ve sloupci `Akcie`.

Do sešitu se zapisují jen sloupce Investovano a Hodnota na řádcích akcií. Nic jiného
se nemění. Sešit má grafy, tabulky a webextension, které by openpyxl při uložení
poškodil, takže se hodnoty zapisují přímo do XML uvnitř .xlsx.

Spuštění z kořene repa:
    .\\.venv\\Scripts\\python.exe projects/stocks-tracker/scripts/xtb_sync.py --dry-run
    .\\.venv\\Scripts\\python.exe projects/stocks-tracker/scripts/xtb_sync.py --write
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
import unicodedata
import warnings
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from shared.lib.repo import data_dir, dated  # noqa: E402

PROJECT = "stocks-tracker"

ROOT = Path(r"D:\Work\Finance\Test")  # kořen se složkami vlastníků
SPORENI = ROOT / "sporeni.xlsx"

CURRENCIES = ("USD", "EUR", "CZK", "GBP")

# Sloupce, bez kterých list akcií nejde zpracovat. `Nazev akcie` (s tickerem) je
# volitelný — bez něj se páruje jen podle názvu.
REQUIRED = ("Akcie", "Vlastnik", "Broker", "Mena", "Investovano", "Hodnota")
TICKER_COL = "Nazev akcie"

# Názvy, které se mezi sešitem a exportem liší víc, než srovná normalizace.
# Uplatní se jen u řádků bez tickeru.
ALIASES = {
    "treasurybondib01usd": "treasurybond01y",
    "arcelormital": "arcelormittal",
    "baesystems": "bae",
    "mcdonald": "mcdonalds",
    "coreeurostoxx50": "eurostoxx50",
}


def _plain(text: str) -> str:
    """Malá písmena bez diakritiky, jen alfanumerické znaky."""
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", text.lower())


def norm(name: str) -> str:
    key = _plain(name)
    return ALIASES.get(key, key)


def norm_owner(name: str) -> str:
    return _plain(name)


def row_ticker(text) -> str:
    """Ticker ze sloupce `Nazev akcie` ve tvaru 'GOOGL.US, Alphabet Inc'."""
    if not text:
        return ""
    first = str(text).split(",", 1)[0].strip().upper()
    return first if re.fullmatch(r"[A-Z0-9]+\.[A-Z]{1,3}", first) else ""


# ---------------------------------------------------------------- exporty


def find_exports(root: Path) -> dict[tuple[str, str], dict[str, Path]]:
    """Projde strom vlastník/broker a vrátí {(vlastnik, broker): {mena: soubor}}.

    Hlídá, že pro jednu měnu je ve složce právě jeden soubor — jinak by nebylo jasné,
    který se má použít. Prázdné složky vlastníků/brokerů se přeskočí.
    """
    if not root.is_dir():
        raise SystemExit(f"Kořen exportů neexistuje: {root}")

    found: dict[tuple[str, str], dict[str, Path]] = {}
    problems: list[str] = []

    for owner_dir in sorted(p for p in root.iterdir() if p.is_dir()):
        for broker_dir in sorted(p for p in owner_dir.iterdir() if p.is_dir()):
            files = [
                f
                for f in sorted(broker_dir.glob("*.xlsx"))
                if not f.name.startswith("~$") and not f.name.endswith(".bak.xlsx")
            ]
            by_currency: dict[str, list[Path]] = {}
            for f in files:
                currency = f.name.split("_", 1)[0].upper()
                if currency not in CURRENCIES:
                    problems.append(
                        f"{owner_dir.name}/{broker_dir.name}/{f.name}: "
                        f"název nezačíná měnou ({', '.join(CURRENCIES)})"
                    )
                    continue
                by_currency.setdefault(currency, []).append(f)

            for currency, group in by_currency.items():
                if len(group) > 1:
                    names = ", ".join(f.name for f in group)
                    problems.append(
                        f"{owner_dir.name}/{broker_dir.name}: pro {currency} je souborů "
                        f"{len(group)} ({names}) — nech ve složce jen ten aktuální"
                    )

            usable = {c: g[0] for c, g in by_currency.items() if len(g) == 1}
            if usable:
                found[(owner_dir.name, broker_dir.name)] = usable

    if problems:
        raise SystemExit("Exporty nejsou v pořádku:\n  " + "\n  ".join(problems))
    if not found:
        raise SystemExit(f"Pod {root} nejsou žádné exporty ve tvaru <Vlastnik>/<Broker>/<MENA>_*.xlsx")
    return found


def read_export(path: Path) -> dict:
    """Souhrnné řádky exportu XTB (ty s vyplněnou kategorií) + celková hodnota z hlavičky.

    Investováno = Value - Net Profit, tj. hodnota bez zisku. Sedí na Volume * Open price.
    Vrací {"positions": [...], "total_value": float | None}.
    """
    # Bez read_only: export nemá v XML rozsah listu a v režimu jen pro čtení
    # by openpyxl viděl prázdný list. Export XTB nemá výchozí styl — varování je neškodné.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        wb = openpyxl.load_workbook(path, data_only=True)
    if "Open Positions" not in wb.sheetnames:
        raise SystemExit(f"{path.name}: chybí list 'Open Positions' — je to export otevřených pozic?")
    ws = wb["Open Positions"]

    total_value = None
    header_row = None
    for i, row in enumerate(ws.iter_rows(max_col=4, values_only=True), start=1):
        if row[1] == "Value" and isinstance(row[2], (int, float)):
            total_value = float(row[2])
        if row[1] == "Instrument/Position":
            header_row = i
            break
    if header_row is None:
        raise SystemExit(f"{path.name}: na listu 'Open Positions' chybí hlavička 'Instrument/Position'")

    positions, seen = [], set()
    for row in ws.iter_rows(min_row=header_row + 1, max_col=14, values_only=True):
        name, ticker, category = row[1], row[2], row[3]
        value, net = row[6], row[13]
        if not category or not name:  # prázdná kategorie = dílčí obchod, ne souhrn
            continue
        ticker = str(ticker or "").strip().upper()
        if ticker in seen:
            raise SystemExit(f"{path.name}: dva souhrnné řádky pro {ticker} — mapování by bylo nejisté.")
        seen.add(ticker)
        positions.append(
            {
                "name": str(name),
                "ticker": ticker,
                "invested": (value or 0) - (net or 0),
                "value": value or 0,
            }
        )
    wb.close()
    return {"positions": positions, "total_value": total_value}


# ---------------------------------------------------------------- sešit


class SheetLayout:
    """Kde na listu leží data — sloupce se hledají podle hlaviček, ne natvrdo."""

    def __init__(self, ws) -> None:
        self.title = ws.title
        self.headers = {
            str(c.value).strip(): c.column_letter for c in ws[1] if c.value is not None
        }
        missing = [h for h in REQUIRED if h not in self.headers]
        if missing:
            raise SystemExit(f"List '{ws.title}': chybí sloupce {', '.join(missing)}")
        self.first_row, self.last_row = self._data_rows(ws)

    @staticmethod
    def _data_rows(ws) -> tuple[int, int]:
        """Rozsah řádků z tabulky, která začíná v A1 — ne z max_row (pod tabulkou jsou součty)."""
        for name in ws.tables:
            start, end = ws.tables[name].ref.split(":")
            if start == "A1":
                return 2, int(re.sub(r"[A-Z]", "", end))
        raise SystemExit(f"List '{ws.title}': nenašel jsem tabulku začínající v A1")

    def get(self, ws, header: str, row: int):
        col = self.headers.get(header)
        return ws[f"{col}{row}"].value if col else None

    def cell(self, header: str, row: int) -> str:
        return f"{self.headers[header]}{row}"


def find_sheet(wb) -> str:
    """List akcií = jediný list, jehož první řádek má všechny sloupce z REQUIRED."""
    hits = []
    for title in wb.sheetnames:
        headers = {str(c.value).strip() for c in wb[title][1] if c.value is not None}
        if set(REQUIRED) <= headers:
            hits.append(title)
    if len(hits) != 1:
        raise SystemExit(
            f"Čekám právě jeden list se sloupci {', '.join(REQUIRED)}; nalezeno: {hits or 'žádný'}"
        )
    return hits[0]


def plan(sporeni: Path, root: Path) -> tuple[dict[str, dict[str, float]], list[dict]]:
    """Vrátí (buňky k zápisu podle listů, přehled po vlastníkovi a brokerovi)."""
    exports = find_exports(root)
    wb = openpyxl.load_workbook(sporeni)
    title = find_sheet(wb)
    ws = wb[title]
    layout = SheetLayout(ws)

    cells: dict[str, dict[str, float]] = {}
    blocks: list[dict] = []

    for (owner, broker), files in sorted(exports.items()):
        exp = {ccy: read_export(path) for ccy, path in files.items()}
        by_ticker = {
            ccy: {p["ticker"]: p for p in e["positions"] if p["ticker"]} for ccy, e in exp.items()
        }
        by_name = {ccy: {norm(p["name"]): p for p in e["positions"]} for ccy, e in exp.items()}
        claimed: dict[tuple[str, str], int] = {}  # (mena, ticker) -> řádek

        rows, unmatched, conflicts, problems = [], [], [], []
        all_rows = []  # všechny řádky vlastníka+brokera, i s jinou měnou — pro nápovědy

        for row in range(layout.first_row, layout.last_row + 1):
            name = layout.get(ws, "Akcie", row)
            if not name:
                continue
            if norm_owner(layout.get(ws, "Vlastnik", row) or "") != norm_owner(owner):
                continue
            if _plain(layout.get(ws, "Broker", row) or "") != _plain(broker):
                continue

            ccy = str(layout.get(ws, "Mena", row) or "").strip().upper()
            ticker = row_ticker(layout.get(ws, TICKER_COL, row))
            all_rows.append({"row": row, "name": str(name), "ticker": ticker, "ccy": ccy})

            if ccy not in CURRENCIES:
                problems.append(f"ř. {row} — {name}: neznámá měna ve sloupci Mena ('{ccy}'), řádek přeskočen")
                continue

            pos, how = None, ""
            if ccy in exp:
                if ticker:
                    pos, how = by_ticker[ccy].get(ticker), "ticker"
                else:
                    pos, how = by_name[ccy].get(norm(name)), "název"

            if pos is not None:
                key = (ccy, pos["ticker"])
                if key in claimed:
                    conflicts.append(
                        f"ř. {row} — {name}: pozice {pos['ticker']} už je přiřazená ř. {claimed[key]}, "
                        "řádek se nezapíše"
                    )
                    continue
                claimed[key] = row

            inv_ref, val_ref = layout.cell("Investovano", row), layout.cell("Hodnota", row)
            if pos is None:
                # Řádek bez pozice v exportu se jen nahlásí — hodnoty zůstanou, jak jsou.
                unmatched.append(
                    {
                        "row": row,
                        "name": str(name),
                        "ticker": ticker,
                        "ccy": ccy,
                        "inv": ws[inv_ref].value,
                        "val": ws[val_ref].value,
                        "reason": f"pro {ccy} není export" if ccy not in exp else "v exportu není otevřená pozice",
                    }
                )
                continue

            new_inv, new_val = round(pos["invested"]), round(pos["value"])
            entry = {
                "row": row,
                "name": str(name),
                "ticker": ticker,
                "export_ticker": pos["ticker"],
                "ccy": ccy,
                "how": how,
                "old_inv": ws[inv_ref].value,
                "new_inv": new_inv,
                "old_val": ws[val_ref].value,
                "new_val": new_val,
            }
            entry["changed"] = (entry["old_inv"] or 0, entry["old_val"] or 0) != (new_inv, new_val) or (
                entry["old_inv"] is None or entry["old_val"] is None
            )
            rows.append(entry)
            cells.setdefault(title, {})[inv_ref] = new_inv
            cells.setdefault(title, {})[val_ref] = new_val

        leftover = []
        for ccy, e in exp.items():
            for p in e["positions"]:
                if (ccy, p["ticker"]) in claimed:
                    continue
                leftover.append({**p, "ccy": ccy, "hint": _leftover_hint(p, ccy, all_rows)})

        checks = []
        for ccy, e in sorted(exp.items()):
            export_sum = sum(p["value"] for p in e["positions"])
            mapped = sum(r["new_val"] for r in rows if r["ccy"] == ccy)
            checks.append(
                {"ccy": ccy, "header": e["total_value"], "export_sum": export_sum, "mapped": mapped}
            )

        blocks.append(
            {
                "owner": owner,
                "broker": broker,
                "sheet": title,
                "files": files,
                "rows": rows,
                "unmatched": unmatched,
                "leftover": leftover,
                "conflicts": conflicts,
                "problems": problems,
                "checks": checks,
            }
        )

    return cells, blocks


def _leftover_hint(pos: dict, ccy: str, all_rows: list[dict]) -> str:
    """Proč se pozice z exportu nepřiřadila — hledá řádek, který ji nejspíš myslí."""
    for r in all_rows:
        if pos["ticker"] and r["ticker"] == pos["ticker"] and r["ccy"] != ccy:
            return f"ř. {r['row']} {r['name']} má stejný ticker, ale Mena={r['ccy'] or '—'} (export je {ccy})"
    for r in all_rows:
        if norm(r["name"]) == norm(pos["name"]):
            if r["ccy"] != ccy:
                return f"ř. {r['row']} {r['name']} má Mena={r['ccy'] or '—'} (export je {ccy})"
            if r["ticker"] and r["ticker"] != pos["ticker"]:
                return f"ř. {r['row']} {r['name']} má ticker {r['ticker']}, export {pos['ticker']}"
    return "v sešitu žádný podobný řádek"


# ---------------------------------------------------------------- report


def _fmt(v) -> str:
    return "—" if v is None else f"{v}"


def _pair(old, new) -> str:
    if old == new:
        return f"{old}"
    return f"{_fmt(old)} → **{new}**"


def markdown_report(blocks: list[dict], sporeni: Path, root: Path, written: bool) -> str:
    total_changed = sum(len([r for r in b["rows"] if r["changed"]]) for b in blocks)
    lines = [
        f"# Aktualizace {sporeni.name}",
        "",
        f"- Spuštěno: {datetime.now():%Y-%m-%d %H:%M}",
        f"- Kořen exportů: `{root}`",
        f"- Režim: {'**zapsáno do sešitu**' if written else 'jen výpis, nic se nezapsalo'}",
        f"- Změněných řádků celkem: **{total_changed}**",
        "",
        "Zapisují se jen sloupce Investovano a Hodnota. Export `<MENA>_*.xlsx` plní jen",
        "řádky se stejnou hodnotou ve sloupci `Mena`. Řádky bez pozice v exportu se jen",
        "vypíšou — jejich hodnoty zůstanou, jak jsou.",
    ]

    for block in blocks:
        changed = [r for r in block["rows"] if r["changed"]]
        lines += ["", f"## {block['owner']} — {block['broker']}", ""]
        if not block["rows"] and not block["unmatched"]:
            lines += [
                f"**Export existuje, ale na listu `{block['sheet']}` nemá "
                f"{block['owner']} / {block['broker']} jediný řádek — nezapsalo se nic.**",
                "",
            ]
        lines += [
            f"- List: `{block['sheet']}`",
            f"- Změněných řádků: **{len(changed)}** z {len(block['rows'])}",
            "- Exporty: " + ", ".join(f"{c} `{p.name}`" for c, p in sorted(block["files"].items())),
            "",
            "### Kontrola součtů",
            "",
            "| Měna | Hlavička exportu (Value) | Součet pozic v exportu | Zapsáno do sešitu | Nezapsáno |",
            "|---|---:|---:|---:|---:|",
        ]
        for c in block["checks"]:
            lines.append(
                f"| {c['ccy']} | {_fmt(c['header'] and round(c['header']))} | {round(c['export_sum'])} | "
                f"{c['mapped']} | {'**' + str(round(c['export_sum']) - c['mapped']) + '**' if round(c['export_sum']) - c['mapped'] > 1 else '0'} |"
            )

        lines += [
            "",
            "### Řádky",
            "",
            "| Ř. | Akcie | Ticker | Mena | Páruje se | Investováno | Hodnota |",
            "|---:|---|---|---|---|---:|---:|",
        ]
        for r in block["rows"]:
            mark = "" if r["changed"] else " _(beze změny)_"
            lines.append(
                f"| {r['row']} | {r['name']}{mark} | {r['ticker'] or '—'} | {r['ccy']} | "
                f"{r['how'] or '—'} | {_pair(r['old_inv'], r['new_inv'])} | {_pair(r['old_val'], r['new_val'])} |"
            )

        if block["leftover"]:
            lines += [
                "",
                "### POZOR: pozice z exportu, které nemají řádek v sešitu",
                "",
                "Jejich peníze se nikam nezapočítají. Oprav ticker / Mena v řádku, nebo řádek přidej.",
                "",
                "| Měna | Ticker | Název v exportu | Hodnota | Nejspíš |",
                "|---|---|---|---:|---|",
            ]
            for p in block["leftover"]:
                lines.append(
                    f"| {p['ccy']} | {p['ticker']} | {p['name']} | {round(p['value'])} | {p['hint']} |"
                )
        if block["conflicts"]:
            lines += ["", "### POZOR: dva řádky na jednu pozici", ""]
            lines += [f"- {c}" for c in block["conflicts"]]
        if block["problems"]:
            lines += ["", "### Přeskočené řádky", ""]
            lines += [f"- {p}" for p in block["problems"]]
        if block["unmatched"]:
            lines += [
                "",
                "### Řádky bez pozice v exportu (nezměněny)",
                "",
                "| Ř. | Akcie | Ticker | Mena | Investováno | Hodnota | Důvod |",
                "|---:|---|---|---|---:|---:|---|",
            ]
            lines += [
                f"| {r['row']} | {r['name']} | {r['ticker'] or '—'} | {r['ccy']} | "
                f"{_fmt(r['inv'])} | {_fmt(r['val'])} | {r['reason']} |"
                for r in block["unmatched"]
            ]

    return "\n".join(lines) + "\n"


def console_summary(blocks: list[dict]) -> str:
    lines = []
    for block in blocks:
        changed = [r for r in block["rows"] if r["changed"]]
        lines.append(f"\n{block['owner']} / {block['broker']}  →  list '{block['sheet']}'")
        lines.append(
            f"  exporty: {', '.join(sorted(block['files']))}"
            f"   změněných řádků: {len(changed)} z {len(block['rows'])}"
        )
        if not block["rows"] and not block["unmatched"]:
            lines.append(
                f"  !! na listu není jediný řádek s Vlastnik={block['owner']} "
                f"a Broker={block['broker']} — nezapsalo se nic"
            )
        for c in block["checks"]:
            lines.append(
                f"  {c['ccy']}: hodnota v exportu {round(c['export_sum'])}, do sešitu {c['mapped']}"
            )
        for r in changed:
            lines.append(
                f"    ř.{r['row']:>3} {r['name']:<28} {r['ccy']} "
                f"inv {_fmt(r['old_inv'])}->{r['new_inv']}, hodn {_fmt(r['old_val'])}->{r['new_val']}"
            )
        for r in block["unmatched"]:
            lines.append(
                f"  - ř.{r['row']:>3} {r['name']:<28} {r['ccy']} bez pozice v exportu "
                f"({r['reason']}) — ponecháno inv {_fmt(r['inv'])}, hodn {_fmt(r['val'])}"
            )
        for p in block["leftover"]:
            lines.append(f"  ! bez řádku: {p['ccy']} {p['ticker']} ({p['name']}, {round(p['value'])}) — {p['hint']}")
        for c in block["conflicts"]:
            lines.append(f"  ! {c}")
        for p in block["problems"]:
            lines.append(f"  ! {p}")
    return "\n".join(lines)


# ---------------------------------------------------------------- zápis


def sheet_paths(path: Path) -> dict[str, str]:
    """{název listu: cesta k jeho XML v zipu}."""
    with zipfile.ZipFile(path) as z:
        wb = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    out = {}
    for name, rid in re.findall(r'<sheet name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        target = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"', rels)
        if target:
            out[name.replace("&amp;", "&")] = "xl/" + target.group(1).lstrip("/").removeprefix("xl/")
    return out


def write_cells(path: Path, per_sheet: dict[str, dict[str, float]]) -> Path:
    """Přepíše hodnoty daných buněk přímo v XML; zbytek sešitu zůstane beze změny."""
    lock = path.with_name("~$" + path.name)
    if lock.exists():
        raise SystemExit(
            f"Sešit je otevřený v Excelu ({lock.name}). Zavři ho — jinak by Excel při "
            f"uložení přepsal, co skript zapíše."
        )
    backup = path.with_suffix(".bak.xlsx")
    shutil.copy2(path, backup)

    with zipfile.ZipFile(path) as zin:
        items = {n: zin.read(n) for n in zin.namelist()}
        infos = {i.filename: i for i in zin.infolist()}

    paths = sheet_paths(path)
    for title, cells in per_sheet.items():
        xml_name = paths.get(title)
        if xml_name is None:
            raise SystemExit(f"List '{title}' nemá v sešitu XML — struktura se změnila.")
        xml = items[xml_name].decode("utf-8")
        for ref, value in cells.items():
            pattern = re.compile(rf'<c r="{ref}"([^>]*?)(/>|>.*?</c>)', re.DOTALL)
            match = pattern.search(xml)
            if not match:
                raise SystemExit(f"Buňka {ref} na listu '{title}' v XML nenalezena.")
            attrs = re.sub(r'\st="[^"]*"', "", match.group(1))  # číslo nesmí mít typ textu
            attrs = attrs.rstrip("/").rstrip()
            xml = xml[: match.start()] + f'<c r="{ref}"{attrs}><v>{value}</v></c>' + xml[match.end() :]
        items[xml_name] = xml.encode("utf-8")

    # Excel má po otevření přepočítat vzorce — jinak by ukazoval staré cache hodnoty.
    wbxml = items["xl/workbook.xml"].decode("utf-8")
    if "<calcPr" in wbxml:
        wbxml = re.sub(r'\sfullCalcOnLoad="[^"]*"', "", wbxml)
        wbxml = re.sub(r"<calcPr([^>]*?)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', wbxml, count=1)
    else:
        wbxml = wbxml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
    items["xl/workbook.xml"] = wbxml.encode("utf-8")

    # calcChain je po změně neplatný; Excel si ho postaví znovu. Musí zmizet i odkazy
    # na něj, jinak Excel hlásí poškozený soubor.
    if items.pop("xl/calcChain.xml", None) is not None:
        for name in ("[Content_Types].xml", "xl/_rels/workbook.xml.rels"):
            text = items[name].decode("utf-8")
            text = re.sub(r"<Override[^>]*calcChain\.xml\"[^>]*/>", "", text)
            text = re.sub(r"<Relationship[^>]*calcChain\.xml\"[^>]*/>", "", text)
            items[name] = text.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items.items():
            zout.writestr(zipfile.ZipInfo(name, date_time=infos[name].date_time), data)
    return backup


# ---------------------------------------------------------------- main


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--write", action="store_true", help="zapsat do sešitu (jinak jen výpis)")
    ap.add_argument("--dry-run", action="store_true", help="jen výpis (výchozí chování)")
    ap.add_argument("--root", type=Path, default=ROOT, help="kořen se složkami vlastníků")
    ap.add_argument("--sporeni", type=Path, default=None, help="cesta k sešitu (výchozí <root>/sporeni.xlsx)")
    ap.add_argument("--no-report", action="store_true", help="neukládat markdown report")
    args = ap.parse_args()

    # Konzole je na Windows cp1250 a české znaky by ji shodily.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    sporeni = args.sporeni or (args.root / "sporeni.xlsx")
    cells, blocks = plan(sporeni, args.root)

    print(console_summary(blocks))

    if args.write:
        backup = write_cells(sporeni, cells)
        print(f"\nZapsáno do {sporeni}")
        print(f"Záloha původního souboru: {backup}")
    else:
        print("\n(jen výpis — do sešitu se nic nezapsalo, zápis provede --write)")

    if not args.no_report:
        suffix = "" if args.write else "_dry-run"
        out = data_dir(PROJECT) / dated(f"xtb_sync{suffix}", ".md")
        out.write_text(markdown_report(blocks, sporeni, args.root, args.write), encoding="utf-8")
        print(f"Report: {out}")


if __name__ == "__main__":
    main()
