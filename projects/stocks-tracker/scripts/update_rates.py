"""Doplní do sporeni.xlsx aktuální kurzy měn k CZK podle denního kurzu ČNB.

Kurzy leží v pojmenované oblasti `Kurzy` na listu Akcie (dnes `A60:B63`):
sloupec A = kód měny, sloupec B = kolik CZK za 1 jednotku. Skript přepíše jen
sloupec B, a to u měn, které ČNB vyhlašuje. CZK zůstává 1.

Zápis jde přímo do XML uvnitř .xlsx (stejně jako `xtb_sync.py`), aby se nepoškodily
grafy a tabulky.

Spuštění z kořene repa:
    .\\.venv\\Scripts\\python.exe projects/stocks-tracker/scripts/update_rates.py
    .\\.venv\\Scripts\\python.exe projects/stocks-tracker/scripts/update_rates.py --write
    .\\.venv\\Scripts\\python.exe projects/stocks-tracker/scripts/update_rates.py --date 2026-09-01 --write
"""

from __future__ import annotations

import argparse
import re
import sys
import urllib.request
from datetime import date, datetime
from pathlib import Path

import openpyxl

from xtb_sync import ROOT, write_cells  # stejný projekt, stejná složka

CNB_URL = (
    "https://www.cnb.cz/cs/financni-trhy/devizovy-trh/kurzy-devizoveho-trhu/"
    "kurzy-devizoveho-trhu/denni_kurz.txt?date={date}"
)
RANGE_NAME = "Kurzy"


def fetch_cnb(on: date) -> tuple[date, dict[str, float]]:
    """Denní kurz ČNB: (datum vyhlášení, {kód: CZK za 1 jednotku}).

    Pro víkend/svátek vrací ČNB poslední platný kurz — skutečné datum je v hlavičce.
    """
    url = CNB_URL.format(date=on.strftime("%d.%m.%Y"))
    with urllib.request.urlopen(url, timeout=20) as resp:
        text = resp.read().decode("utf-8")
    lines = text.strip().splitlines()
    valid = datetime.strptime(lines[0].split()[0], "%d.%m.%Y").date()
    rates = {}
    for line in lines[2:]:
        _country, _name, amount, code, rate = line.split("|")
        rates[code] = float(rate.replace(",", ".")) / int(amount)
    return valid, rates


def rate_cells(sporeni: Path) -> tuple[str, list[tuple[str, str, float | None]]]:
    """Najde oblast `Kurzy`: (list, [(kód měny, buňka s kurzem, současná hodnota)])."""
    wb = openpyxl.load_workbook(sporeni)
    for ws in wb.worksheets:
        dn = ws.defined_names.get(RANGE_NAME)
        if dn is not None:
            break
    else:
        dn = wb.defined_names.get(RANGE_NAME)
        if dn is None:
            raise SystemExit(f"V sešitu není pojmenovaná oblast '{RANGE_NAME}'.")

    (title, ref), *_ = list(dn.destinations)
    ws = wb[title]
    start, end = ref.replace("$", "").split(":")
    first, last = int(re.sub(r"\D", "", start)), int(re.sub(r"\D", "", end))
    code_col = re.sub(r"\d", "", start)
    rate_col = re.sub(r"\d", "", end)

    out = []
    for row in range(first, last + 1):
        code = ws[f"{code_col}{row}"].value
        if code:
            out.append((str(code).strip().upper(), f"{rate_col}{row}", ws[f"{rate_col}{row}"].value))
    return title, out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--write", action="store_true", help="zapsat do sešitu (jinak jen výpis)")
    ap.add_argument("--date", type=date.fromisoformat, default=date.today(), help="datum kurzu YYYY-MM-DD")
    ap.add_argument("--sporeni", type=Path, default=ROOT / "sporeni.xlsx", help="cesta k sešitu")
    args = ap.parse_args()

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    valid, rates = fetch_cnb(args.date)
    title, cells = rate_cells(args.sporeni)

    print(f"Kurz ČNB vyhlášený {valid.isoformat()} → list '{title}', oblast {RANGE_NAME}")
    updates: dict[str, float] = {}
    for code, ref, old in cells:
        if code == "CZK":
            print(f"  {ref}  CZK  {old}  (beze změny)")
            continue
        if code not in rates:
            print(f"  {ref}  {code}  {old}  !! ČNB tuto měnu nevyhlašuje, ponecháno")
            continue
        new = round(rates[code], 4)
        updates[ref] = new
        print(f"  {ref}  {code}  {old} → {new}")

    if not args.write:
        print("\n(jen výpis — do sešitu se nic nezapsalo, zápis provede --write)")
        return
    if not updates:
        print("\nNení co zapsat.")
        return
    backup = write_cells(args.sporeni, {title: updates})
    print(f"\nZapsáno do {args.sporeni}")
    print(f"Záloha původního souboru: {backup}")


if __name__ == "__main__":
    main()
